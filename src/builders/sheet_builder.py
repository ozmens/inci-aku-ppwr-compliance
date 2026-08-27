"""Per-sheet creation: headers, Excel Table, no formulas/formatting."""

from __future__ import annotations

from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet

from models.registry import SchemaRegistry
from models.table_definition import TableDefinition

# Placeholder data row so ListObject + named PK ranges have a body cell.
PLACEHOLDER_ROWS = 1


def excel_col_letter(index: int) -> str:
    """1-based column index → Excel column letter(s)."""
    result = []
    n = index
    while n > 0:
        n, rem = divmod(n - 1, 26)
        result.append(chr(65 + rem))
    return "".join(reversed(result))


class SheetBuilder:
    """Writes one normalized worksheet per TableDefinition."""

    def __init__(self, registry: SchemaRegistry) -> None:
        self.registry = registry

    def table_display_name(self, table: TableDefinition) -> str:
        return f"T_{table.name}"

    def write_sheet(self, ws: Worksheet, table: TableDefinition) -> Table:
        headers = list(table.column_names)
        if not headers:
            raise ValueError(f"{table.name} has no columns")

        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # Empty placeholder body rows (no sample business data)
        for row_idx in range(2, 2 + PLACEHOLDER_ROWS):
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx, value=None)

        last_col = excel_col_letter(len(headers))
        last_row = 1 + PLACEHOLDER_ROWS
        ref = f"A1:{last_col}{last_row}"

        excel_table = Table(displayName=self.table_display_name(table), ref=ref)
        # No TableStyleInfo → no visual formatting
        ws.add_table(excel_table)

        # Freeze panes / column widths intentionally omitted (no formatting)
        return excel_table

    def pk_column_index(self, table: TableDefinition) -> int:
        try:
            return list(table.column_names).index(table.primary_key) + 1
        except ValueError as exc:
            raise ValueError(
                f"PK {table.primary_key} missing on {table.name}"
            ) from exc

    def pk_range_formula(self, table: TableDefinition) -> str:
        """Sheet-local absolute range for PK body cells (named range target)."""
        col = excel_col_letter(self.pk_column_index(table))
        # Expandable placeholder range for data-validation lists
        return f"'{table.name}'!${col}$2:${col}$1048576"

    def column_letter(self, table: TableDefinition, column_name: str) -> str:
        idx = list(table.column_names).index(column_name) + 1
        return excel_col_letter(idx)
